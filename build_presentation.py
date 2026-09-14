"""
Builds the Model Review Committee presentation as a real .pptx file.

Audience: mostly business, a few technical. Every slide leads with money and
plain language; the technical evidence sits underneath for the people who
want it. Light analyst-report styling, real notebook-generated charts.

Usage:
    python build_presentation.py
Output:
    reports/Credit_Scoring_Model_Review.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

# ---------------------------------------------------------------- theme ----

BG = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF3, 0xF6, 0xFA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
CARD_TINT = RGBColor(0xF6, 0xF9, 0xFB)
BORDER = RGBColor(0xDD, 0xE4, 0xEC)
NAVY = RGBColor(0x14, 0x22, 0x36)
MUTED = RGBColor(0x5B, 0x6B, 0x7E)
FAINT = RGBColor(0x93, 0xA1, 0xB2)
BLUE = RGBColor(0x0B, 0x5F, 0xA5)
BLUE_DARK = RGBColor(0x0A, 0x3D, 0x66)
BLUE_TINT = RGBColor(0xEC, 0xF3, 0xF9)
ORANGE = RGBColor(0xE0, 0x6E, 0x1D)
ORANGE_TINT = RGBColor(0xFD, 0xF0, 0xE4)
GREEN = RGBColor(0x1E, 0x8E, 0x5A)
GREEN_TINT = RGBColor(0xE9, 0xF6, 0xEF)
RED = RGBColor(0xB8, 0x3A, 0x2E)
RED_TINT = RGBColor(0xFB, 0xEC, 0xEA)
AMBER = RGBColor(0xB5, 0x7C, 0x0E)
AMBER_TINT = RGBColor(0xFB, 0xF3, 0xE1)
DARK_BG = RGBColor(0x0A, 0x0E, 0x16)

HEAD_FONT = "Segoe UI Semibold"
BODY_FONT = "Segoe UI"
NUM_FONT = "Segoe UI Semibold"
MONO_FONT = "Consolas"

SW, SH = Inches(13.333), Inches(7.5)
FIG = "reports/figures"
TOTAL = 13

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def _no_autofit(tf):
    el = tf._txBody
    bodyPr = el.find(qn('a:bodyPr'))
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        e = bodyPr.find(qn(tag))
        if e is not None:
            bodyPr.remove(e)
    bodyPr.append(el.makeelement(qn('a:noAutofit'), {}))


def new_slide(page=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Pt(10))
    top.fill.solid(); top.fill.fore_color.rgb = BLUE_DARK
    top.line.fill.background(); top.shadow.inherit = False
    tick = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Pt(10), Inches(2.6), Pt(4))
    tick.fill.solid(); tick.fill.fore_color.rgb = ORANGE
    tick.line.fill.background(); tick.shadow.inherit = False
    if page is not None:
        tb = s.shapes.add_textbox(Inches(12.05), Inches(0.24), Inches(1.05), Inches(0.35))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = f"{page:02d} / {TOTAL}"
        r.font.name = BODY_FONT; r.font.size = Pt(11); r.font.color.rgb = FAINT
    return s


def pill(s, text, left, top, color=BLUE_DARK, text_color=RGBColor(0xFF, 0xFF, 0xFF),
         width=None, size=11.5):
    w = width if width else Inches(0.4 + 0.0145 * size * len(text))
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), w, Inches(0.36))
    try:
        sh.adjustments[0] = 0.5
    except Exception:
        pass
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = Inches(0.14); tf.margin_right = Inches(0.14)
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text.upper()
    r.font.name = BODY_FONT; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = text_color
    try:
        r._r.get_or_add_rPr().set('spc', '80')
    except Exception:
        pass
    return sh


def kicker(s, text, top=0.42, color=BLUE_DARK):
    return pill(s, text, 0.68, top, color=color)


def headline(s, text, top=1.0, size=28, width=11.9, color=NAVY):
    tb = s.shapes.add_textbox(Inches(0.68), Inches(top), Inches(width), Inches(1.3))
    tf = tb.text_frame; tf.word_wrap = True; _no_autofit(tf)
    r = tf.paragraphs[0].add_run(); r.text = text
    r.font.name = HEAD_FONT; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def body_text(s, text, left, top, width, height=0.6, size=13, color=MUTED,
              align=PP_ALIGN.LEFT, italic=False, font=BODY_FONT, bold=False):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True; _no_autofit(tf)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.color.rgb = color
    r.font.italic = italic; r.font.bold = bold
    return tb


def card(s, left, top, width, height, line_color=BORDER, fill=CARD):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                            Inches(width), Inches(height))
    try:
        sh.adjustments[0] = 0.04
    except Exception:
        pass
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line_color; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def accent_edge(s, left, top, height, color, width_in=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                            Inches(width_in), Inches(height))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def big_num(s, text, left, top, width, size=44, color=NAVY, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(size / 52 + 0.35))
    tf = tb.text_frame; tf.word_wrap = True; _no_autofit(tf)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = NUM_FONT; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def takeaway(s, left, top, width, height, text, label="KEY TAKEAWAY",
             color=ORANGE, tint=ORANGE_TINT, size=11.5):
    card(s, left, top, width, height, fill=tint, line_color=color)
    accent_edge(s, left, top, height, color)
    body_text(s, label, left + 0.26, top + 0.13, width - 0.5, 0.28, size=10, color=color, bold=True)
    body_text(s, text, left + 0.26, top + 0.42, width - 0.5, height - 0.5, size=size, color=NAVY)


def chevron(s, left, top, size, color=ORANGE):
    sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(left), Inches(top),
                            Inches(size), Inches(size * 0.9))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def image_card(s, path, left, top, max_w, max_h, pad=0.2):
    im = Image.open(path)
    ar = im.size[0] / im.size[1]
    bw, bh = max_w - 2 * pad, max_h - 2 * pad
    if bw / bh > ar:
        h = bh; w = h * ar
    else:
        w = bw; h = w / ar
    card(s, left, top, max_w, max_h, fill=CARD, line_color=BORDER)
    s.shapes.add_picture(path, Inches(left + (max_w - w) / 2), Inches(top + (max_h - h) / 2),
                         width=Inches(w), height=Inches(h))


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def set_cell(cell, text, size=13, color=NAVY, bold=False, font=BODY_FONT,
             align=PP_ALIGN.LEFT, fill=None):
    cell.fill.solid(); cell.fill.fore_color.rgb = fill if fill else CARD
    cell.margin_left = Inches(0.14); cell.margin_right = Inches(0.14)
    cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def ruler(s, left, top, width, zones, marker_pct, marker_label):
    """A benchmark scale: coloured zones with a labelled marker on top."""
    h = 0.42
    x = left
    for frac, col, lab in zones:
        seg_w = width * frac
        seg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(top), Inches(seg_w), Inches(h))
        seg.fill.solid(); seg.fill.fore_color.rgb = col
        seg.line.fill.background(); seg.shadow.inherit = False
        body_text(s, lab, x, top + h + 0.06, seg_w, 0.5, size=10, color=MUTED, align=PP_ALIGN.CENTER)
        x += seg_w
    mx = left + width * marker_pct
    mark = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(mx - 0.02), Inches(top - 0.14),
                              Inches(0.04), Inches(h + 0.28))
    mark.fill.solid(); mark.fill.fore_color.rgb = NAVY
    mark.line.fill.background(); mark.shadow.inherit = False
    body_text(s, marker_label, mx - 1.3, top - 0.45, 2.6, 0.4, size=12.5,
              color=NAVY, bold=True, align=PP_ALIGN.CENTER)


# ============================================================== SLIDE 1 ====
# Title. No numbers, by request: this slide sets the room, it does not brief it.
s = new_slide()
kicker(s, "Consumer Lending  ·  Model Review Committee", top=2.5)
headline(s, "Pricing the Bet", top=2.95, size=54, width=11.5)
body_text(s, "A credit scoring model for the application desk, and what it is worth to the book",
          0.7, 4.35, 11.0, 0.6, size=17, color=MUTED)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(5.1), Inches(1.4), Pt(3))
line.fill.solid(); line.fill.fore_color.rgb = ORANGE
line.line.fill.background(); line.shadow.inherit = False
body_text(s, "Kaustubh Patil", 0.7, 5.3, 6, 0.4, size=14, color=NAVY, bold=True)
body_text(s, "20 August 2026", 0.7, 5.68, 6, 0.4, size=13, color=MUTED)
notes(s, "Open with the one line: I built a model that decides who to lend to, I can prove what it "
         "is worth in dollars, and I found and fixed a compliance problem on the way. Recommendation "
         "is deploy, with conditions. Then move to the two applicants.")

# ============================================================== SLIDE 2 ====
# The problem, told as two people and their money.
s = new_slide(2)
kicker(s, "The Problem")
headline(s, "Two applicants walk in. On paper, they look the same.", top=0.95, size=27)

cw, gap = 5.15, 0.55
lx, rx = 0.7, 0.7 + 5.15 + 0.55 + 0.9
card(s, 0.7, 1.85, cw, 2.75, fill=GREEN_TINT, line_color=GREEN)
accent_edge(s, 0.7, 1.85, 2.75, GREEN)
body_text(s, "APPLICANT A", 1.0, 2.05, 4.6, 0.3, size=11, color=GREEN, bold=True)
body_text(s, "Borrows $20,000 and pays every instalment.", 1.0, 2.4, 4.6, 0.55, size=15, color=NAVY)
big_num(s, "+$1,000", 1.0, 3.0, 4.6, size=40, color=GREEN)
body_text(s, "the bank earns its 5% lifetime margin", 1.0, 3.95, 4.6, 0.4, size=12.5, color=MUTED)

card(s, 6.6, 1.85, cw, 2.75, fill=RED_TINT, line_color=RED)
accent_edge(s, 6.6, 1.85, 2.75, RED)
body_text(s, "APPLICANT B", 6.9, 2.05, 4.6, 0.3, size=11, color=RED, bold=True)
body_text(s, "Borrows $20,000 and stops paying after six months.", 6.9, 2.4, 4.6, 0.55, size=15, color=NAVY)
big_num(s, "−$13,000", 6.9, 3.0, 4.6, size=40, color=RED)
body_text(s, "the bank loses 65% of what it lent", 6.9, 3.95, 4.6, 0.4, size=12.5, color=MUTED)

card(s, 0.7, 4.8, 11.9, 0.95, fill=NAVY, line_color=NAVY)
body_text(s, "One Applicant B costs the bank thirteen Applicant A's.", 1.0, 4.98, 11.3, 0.6,
          size=21, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
takeaway(s, 0.7, 5.95, 11.9, 1.15,
         "About 8 of every 100 loans we approve turn out to be a B. I cannot tell them apart by "
         "looking, and the two mistakes do not cost the same, so the entire job is working out "
         "which one is sitting in front of us before any money moves.",
         label="WHY THIS IS HARD", size=13)
notes(s, "Land the asymmetry, not the percentage. A wrongly declined good customer costs me a small "
         "margin. A wrongly approved bad one costs me most of the loan. Thirteen to one. That ratio "
         "is why the cut-off, not the model, is where the money is decided. If asked: 65% loss given "
         "default, 5% lifetime margin, both named constants in the code.")

# ============================================================== SLIDE 3 ====
s = new_slide(3)
kicker(s, "My Approach")
headline(s, "What I actually did, in four steps", top=0.95, size=27)
steps = [
    ("01", "Gathered the evidence", "307,511 past applications, joined to 1.7M records of how those "
     "same people repaid other lenders."),
    ("02", "Cleaned it honestly", "Fixed the data's lies (one field claimed a thousand years of "
     "employment) and locked away a test set I would not touch until the end."),
    ("03", "Built the signals", "153 measures a lender actually cares about: what they earn, what "
     "they owe, whether they have paid people back before."),
    ("04", "Tried three models", "From a simple points checklist to a modern one, judged on the "
     "same evidence, under the same rules."),
]
w, gap, chev = 2.6, 0.5, 0.34
x, top, h = 0.7, 1.85, 3.15
for idx, (num, title, desc) in enumerate(steps):
    card(s, x, top, w, h, fill=CARD_TINT)
    accent_edge(s, x, top, h, BLUE)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.22), Inches(top + 0.2),
                               Inches(0.54), Inches(0.54))
    badge.fill.solid(); badge.fill.fore_color.rgb = BLUE_DARK
    badge.line.fill.background(); badge.shadow.inherit = False
    btf = badge.text_frame; btf.word_wrap = False
    btf.margin_left = Inches(0.02); btf.margin_right = Inches(0.02)
    btf.margin_top = Inches(0.02); btf.margin_bottom = Inches(0.02)
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = btf.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run(); br.text = num
    br.font.name = NUM_FONT; br.font.size = Pt(14); br.font.bold = True
    br.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    body_text(s, title, x + 0.22, top + 0.88, w - 0.42, 0.5, size=13, color=NAVY, bold=True)
    body_text(s, desc, x + 0.22, top + 1.38, w - 0.42, h - 1.55, size=11.5, color=MUTED)
    if idx < 3:
        chevron(s, x + w + (gap - chev) / 2, top + h / 2 - chev * 0.45, chev)
    x += w + gap
takeaway(s, 0.7, 5.3, 11.9, 1.45,
         "The test set was sealed until the very last step, and every cleaning rule was learned from "
         "the training data only. That matters because it is the difference between a number that "
         "holds up in production and a number that only ever looked good in a notebook.",
         label="WHY YOU CAN TRUST THE NUMBER THAT FOLLOWS", size=12.5)
notes(s, "Keep this fast, 45 seconds. The one point a technical reviewer will test: leakage. Every "
         "fitted transform lives inside one pipeline, refitted per cross-validation fold, so the "
         "test set could not leak into training even by accident. Sentinel detail if asked: "
         "DAYS_EMPLOYED = 365243 on 18% of rows, almost all pensioners, fixed to missing plus a flag.")

# ============================================================== SLIDE 4 ====
s = new_slide(4)
kicker(s, "Results  ·  The Scoreboard")
headline(s, "Three ways to make the call", top=0.95, size=27)
models = [
    ("THE CHECKLIST", "Logistic scorecard", "Adds up points, like an underwriter with a rulebook. "
     "Easy to explain, but it can only draw straight lines.", "0.757", BLUE, False),
    ("THE COMMITTEE", "Random forest", "Hundreds of small opinions, each looking at part of the "
     "file, then a majority vote.", "0.758", BLUE, False),
    ("THE APPRENTICE", "Gradient boosting", "Each round studies the last round's mistakes and "
     "corrects them. Slowest to train, sharpest answer.", "0.774", GREEN, True),
]
w, gap = 3.78, 0.28
x, top, h = 0.7, 1.85, 2.45
for tag, name, desc, score, col, winner in models:
    card(s, x, top, w, h, fill=GREEN_TINT if winner else CARD_TINT,
         line_color=GREEN if winner else BORDER)
    accent_edge(s, x, top, h, col)
    body_text(s, tag, x + 0.24, top + 0.16, w - 0.5, 0.3, size=10.5, color=col, bold=True)
    body_text(s, name, x + 0.24, top + 0.46, w - 0.5, 0.35, size=14, color=NAVY, bold=True)
    body_text(s, desc, x + 0.24, top + 0.88, w - 0.5, 1.0, size=11.5, color=MUTED)
    big_num(s, score, x + 0.24, top + 1.85, 2.0, size=25, color=col)
    if winner:
        pill(s, "selected", x + w - 1.35, top + 1.98, color=GREEN, size=9)
    x += w + gap

body_text(s, "That score is AUROC: given any good customer and any bad one, how often does the "
             "model correctly rank the bad one as riskier?",
          0.7, 4.45, 11.9, 0.35, size=12.5, color=NAVY)
ruler(s, 0.7, 5.5, 11.9,
      [(0.28, RGBColor(0xF2, 0xD5, 0xD2), "0.50 to 0.65\ncoin flip to weak"),
       (0.22, RGBColor(0xFA, 0xEC, 0xD5), "0.65 to 0.70\nusable, not good"),
       (0.28, RGBColor(0xDD, 0xEE, 0xE4), "0.70 to 0.78\nindustry normal on this kind of data"),
       (0.22, RGBColor(0xBF, 0xE3, 0xCF), "0.80 and up\nneeds behavioural data I do not have")],
      marker_pct=0.605, marker_label="mine: 0.774")
takeaway(s, 0.7, 6.65, 11.9, 0.7,
         "0.80 is not the target. On application-and-bureau data alone, without any behavioural "
         "history, this sits where a good model should sit.", label="READING THE SCALE HONESTLY",
         size=12.5)
notes(s, "For the business people: 0.774 means if I line up a good customer and a bad one at random, "
         "the model puts the bad one higher three times out of four. For the technical people: sealed "
         "test AUROC 0.774, Gini 0.547, KS 0.416, and cross-validation, validation and test agree "
         "within half a point. Why not deep learning: on 153 tabular features boosted trees still win.")

# ============================================================== SLIDE 5 ====
s = new_slide(5)
kicker(s, "Results  ·  Separation")
headline(s, "How cleanly does it separate good from bad?", top=0.95, size=27)
image_card(s, f"{FIG}/mod_roc_curves.png", 0.7, 1.85, 6.9, 4.35)
stats = [("3 in 4", "times, it ranks the bad customer above the good one", BLUE),
         ("1 in 3", "of every default sits in the riskiest tenth of the book", GREEN)]
y = 1.85
for val, lab, col in stats:
    card(s, 7.9, y, 4.7, 1.35, fill=CARD_TINT)
    accent_edge(s, 7.9, y, 1.35, col)
    big_num(s, val, 8.15, y + 0.16, 4.2, size=30, color=col)
    body_text(s, lab, 8.15, y + 0.82, 4.2, 0.5, size=12, color=MUTED)
    y += 1.5
takeaway(s, 7.9, 4.9, 4.7, 1.3,
         "Risk is concentrated, not smeared evenly across the book. That concentration is what makes "
         "a single cut-off worth drawing at all.", label="WHY THIS MATTERS", size=12.5)
body_text(s, "The blue curve is the model I chose. The higher it sits above the dotted line, "
             "the better it is telling the two apart.", 0.7, 6.35, 6.9, 0.5, size=11.5, color=MUTED)
notes(s, "Do not explain ROC mechanics unless asked. The business translation is the left number: "
         "three times out of four it puts the bad customer above the good one. If a technical "
         "reviewer asks for the working point: 74.5% of defaults caught while flagging about a "
         "third of the book, and that point is chosen by the profit curve, not by me.")

# ============================================================== SLIDE 6 ====
s = new_slide(6)
kicker(s, "Results  ·  Honesty")
headline(s, "When it says 20%, does 20% actually happen?", top=0.95, size=27)
image_card(s, f"{FIG}/mod_calibration.png", 0.7, 1.85, 6.5, 4.35)
body_text(s, "Yes. Decile by decile, predicted risk lines up with what actually happened.",
          7.55, 1.95, 5.05, 0.7, size=16, color=NAVY, bold=True)
card(s, 7.55, 2.85, 5.05, 1.15, fill=GREEN_TINT, line_color=GREEN)
accent_edge(s, 7.55, 2.85, 1.15, GREEN)
body_text(s, "RISKIEST TENTH OF APPLICANTS", 7.8, 3.0, 4.6, 0.3, size=10, color=GREEN, bold=True)
big_num(s, "28% predicted  →  29% actual", 7.8, 3.3, 4.7, size=17, color=NAVY)
body_text(s, "This is the difference between a model that ranks and a model you can price with. "
             "A ranking tells me who is riskier. An honest probability tells me how much to expect "
             "to lose, which is the number the next slide turns into money.",
          7.55, 4.2, 5.05, 1.5, size=12.5, color=MUTED)
takeaway(s, 7.55, 5.5, 5.05, 1.2,
         "I deliberately refused a common shortcut that would have made the model look better on "
         "paper and broken this chart.", label="A CHOICE I MADE", size=12.5)
notes(s, "The shortcut is class weighting. It boosts the rare class and inflates every predicted "
         "probability. It would have improved nothing that matters and destroyed the dollar maths "
         "and the decline reasons, both of which need the probability to be literally true. "
         "Recalibration trigger is in the monitoring plan: observed above 1.25x predicted in two "
         "or more deciles.")

# ============================================================== SLIDE 7 ====
# The money slide. The break-even insight is the centrepiece.
s = new_slide(7)
kicker(s, "The Decision", color=GREEN)
headline(s, "Where I drew the line, and why it is not arbitrary", top=0.95, size=27)

card(s, 0.7, 1.85, 5.9, 2.5, fill=GREEN_TINT, line_color=GREEN)
accent_edge(s, 0.7, 1.85, 2.5, GREEN)
body_text(s, "NET BENEFIT ON THE SEALED TEST BOOK", 1.0, 2.05, 5.3, 0.3, size=11, color=GREEN, bold=True)
big_num(s, "+$578.7M", 0.98, 2.4, 5.5, size=52, color=GREEN)
body_text(s, "versus approving everyone, across 46,127 applications",
          1.0, 3.65, 5.3, 0.5, size=13, color=NAVY)

card(s, 0.7, 4.5, 5.9, 2.25, fill=CARD_TINT)
accent_edge(s, 0.7, 4.5, 2.25, ORANGE)
body_text(s, "WHY 7%, AND NOT SOME OTHER NUMBER", 1.0, 4.66, 5.3, 0.3, size=10.5,
          color=ORANGE, bold=True)
body_text(s, "A good loan earns 5%. A bad one loses 65%.", 1.0, 4.98, 5.3, 0.35, size=13, color=NAVY)
big_num(s, "5%  ÷  (5% + 65%)  =  7.1%", 1.0, 5.32, 5.4, size=19, color=NAVY)
body_text(s, "That is the break-even: past it, a loan stops making money. I swept the real profit "
             "curve over 46,127 loans and it landed on 7%. Two independent routes, the same answer.",
          1.0, 5.85, 5.3, 0.8, size=11.5, color=MUTED)

image_card(s, f"{FIG}/mod_profit_curve.png", 6.85, 1.85, 5.75, 3.5)
sub_w = 2.75
for i, (val, lab) in enumerate([("63.2%", "of applicants approved at that line"),
                                ("74.5%", "of future defaults stopped before funding")]):
    x = 6.85 + i * (sub_w + 0.25)
    card(s, x, 5.55, sub_w, 1.2, fill=CARD_TINT)
    accent_edge(s, x, 5.55, 1.2, BLUE)
    big_num(s, val, x + 0.2, 5.66, 2.4, size=24, color=BLUE)
    body_text(s, lab, x + 0.2, 6.18, 2.4, 0.5, size=11, color=MUTED)
notes(s, "This is the slide the CIO funds, so slow down. The threshold is leadership's dial, not "
         "mine: move it right and you approve more and earn less, for example 10% approves about "
         "75% of applicants and nets around $480M. The model does not choose the risk appetite, it "
         "prices whichever one you choose. If asked where 578.7 comes from: prevented losses of "
         "$972.0M minus $393.3M of margin given up on good customers I declined.")

# ============================================================== SLIDE 8 ====
s = new_slide(8)
kicker(s, "Explainability")
headline(s, "Every “no” comes with a receipt", top=0.95, size=27)

card(s, 0.7, 1.85, 6.2, 4.15, fill=RED_TINT, line_color=RED)
accent_edge(s, 0.7, 1.85, 4.15, RED)
body_text(s, "WHAT THE CUSTOMER RECEIVES", 1.0, 2.03, 5.6, 0.3, size=10.5, color=RED, bold=True)
body_text(s, "“Your application was declined. The main reasons were:”",
          1.0, 2.38, 5.6, 0.5, size=14, color=NAVY, bold=True)
reasons = ["You are using a high share of the credit you already have",
           "Your external credit score is low",
           "You have several active loans with other lenders",
           "Some payments elsewhere are past due"]
ry = 2.95
for i, rtext in enumerate(reasons, 1):
    body_text(s, f"{i}.   {rtext}", 1.1, ry, 5.5, 0.45, size=12.5, color=NAVY)
    ry += 0.45
body_text(s, "Plain language, no jargon, no column names. Required by FCAC and ECOA, generated "
             "automatically for every single decline.", 1.0, 5.0, 5.6, 0.7, size=11.5, color=MUTED)

card(s, 7.25, 1.85, 5.35, 4.15, fill=CARD_TINT)
accent_edge(s, 7.25, 1.85, 4.15, BLUE)
body_text(s, "HOW I KNOW THOSE ARE THE REAL REASONS", 7.55, 2.03, 4.8, 0.35, size=10.5,
          color=BLUE, bold=True)
body_text(s, "The model does not hand me a score and stay silent. It hands me the score broken "
             "into every factor that built it, and those pieces add up to the score exactly.",
          7.55, 2.45, 4.8, 1.1, size=13, color=NAVY)
body_text(s, "So the four reasons on the letter are not a guess at what mattered, and they are "
             "not written by hand afterwards. They are the four largest pieces of that person's "
             "own number.", 7.55, 3.55, 4.8, 1.2, size=13, color=NAVY)
body_text(s, "Technically: SHAP values, which are additive by construction, mapped to consumer "
             "language across all 279 model columns.", 7.55, 4.85, 4.8, 0.8, size=11, color=MUTED)

takeaway(s, 0.7, 6.15, 11.9, 0.95,
         "Compliance here is not a report written after the fact. It falls out of the same maths "
         "that made the decision, which means the reasons can never drift away from the real cause.",
         size=12.5)
notes(s, "If a risk officer pushes: the additivity is the guarantee. The contributions sum exactly "
         "to the prediction, so I cannot show a reason that was not genuinely one of the drivers. "
         "This is the same applicant I score live in the demo. Global picture if asked: external "
         "scores dominate, then age and tenure, then bureau burden, all in the expected direction.")

# ============================================================== SLIDE 9 ====
s = new_slide(9)
kicker(s, "Risks & Limits", color=AMBER)
headline(s, "Three things that could go wrong, and what I did about each", top=0.95, size=26)
rows = [
    ("The economy turns and everyone gets riskier at once",
     "I built a tripwire that watches for the population drifting, and I test-fired it: a simulated "
     "severe recession trips it. It runs monthly in production.", "TESTED, NOT PROMISED"),
    ("I cannot check the model against a different point in time",
     "This data carries no dates, so an out-of-time test is impossible today. I retrain monthly and "
     "monitor drift instead of pretending the problem is not there.", "COMPENSATING CONTROL"),
    ("The model has only ever seen people we already said yes to",
     "It cannot learn from applicants the old process turned away. Thin files go to a human "
     "reviewer, and correcting this is the first item on my roadmap.", "HUMAN IN THE LOOP"),
]
y = 1.85
TAG_SIZE = 9
for risk, control, tag in rows:
    card(s, 0.7, y, 11.9, 1.4, fill=CARD_TINT)
    accent_edge(s, 0.7, y, 1.4, AMBER)
    body_text(s, risk, 1.0, y + 0.2, 4.2, 0.9, size=13.5, color=NAVY, bold=True)
    body_text(s, control, 5.4, y + 0.2, 4.15, 1.05, size=12, color=MUTED)
    tag_w = 0.4 + 0.0145 * TAG_SIZE * len(tag)
    pill(s, tag, 12.35 - tag_w, y + 0.52, color=AMBER, size=TAG_SIZE)
    y += 1.52
takeaway(s, 0.7, 6.5, 11.9, 0.8,
         "I would rather bring you a model whose blind spots are named and instrumented than one "
         "that looks flawless.", size=12.5)
notes(s, "Drift metric is PSI: below 0.10 stable, 0.10 to 0.25 monitor, above 0.25 investigate. The "
         "simulated severe recession reads 0.34 and the random-split control reads about zero, so "
         "both directions of the control are proven. Segment check: every major segment sits near "
         "the overall line, thin files are the weakest, which is exactly why they get a human.")

# ============================================================== SLIDE 10 ===
s = new_slide(10)
kicker(s, "Fairness & Recommendation")
headline(s, "I found a compliance defect, fixed it, and priced the fix", top=0.95, size=27)
card(s, 0.7, 1.85, 5.75, 2.4, fill=RED_TINT, line_color=RED)
accent_edge(s, 0.7, 1.85, 2.4, RED)
body_text(s, "THE DEFECT I FOUND", 1.0, 2.03, 5.2, 0.3, size=11, color=RED, bold=True)
body_text(s, "The model was using gender as an input. That is prohibited in credit decisions, "
             "no matter how much accuracy it buys.", 1.0, 2.4, 5.2, 0.9, size=13.5, color=NAVY)
body_text(s, "At my cut-off, men were approved at 0.75 of the rate women were, against a 0.80 "
             "regulatory screen. Young applicants sat at 0.62.", 1.0, 3.25, 5.2, 0.9,
          size=12.5, color=MUTED)

card(s, 6.85, 1.85, 5.75, 2.4, fill=GREEN_TINT, line_color=GREEN)
accent_edge(s, 6.85, 1.85, 2.4, GREEN)
body_text(s, "WHAT THE FIX COST", 7.15, 2.03, 5.2, 0.3, size=11, color=GREEN, bold=True)
big_num(s, "0.0005", 7.13, 2.35, 5.2, size=38, color=GREEN)
body_text(s, "of accuracy. I retrained the identical model without gender and lost half a "
             "thousandth of a point. Essentially nothing.", 7.15, 3.05, 5.2, 0.9, size=13, color=NAVY)
body_text(s, "Age effects remain, and they are defensible: older applicants are favoured, "
             "not penalised.", 7.15, 3.72, 5.2, 0.5, size=12, color=MUTED)

card(s, 0.7, 4.45, 11.9, 1.55, fill=BLUE_DARK, line_color=BLUE_DARK)
body_text(s, "RECOMMENDATION:  DEPLOY, WITH CONDITIONS", 1.0, 4.63, 11.3, 0.4, size=17,
          color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
body_text(s, "Ship the gender-free model.   Reasons on every decline.   Monthly drift monitoring.   "
             "Quarterly fairness re-audit.   A human reviewer for thin files.",
          1.0, 5.1, 11.3, 0.75, size=13, color=RGBColor(0xDC, 0xE7, 0xF2))
body_text(s, "A model that beats the baseline, prices its own decisions, explains itself, and "
             "survived a fairness audit. That is what deployable means.",
          0.7, 6.2, 11.9, 0.5, size=13, color=MUTED, italic=True)
notes(s, "Own this, do not apologise for it. Finding it yourself is the strongest thing on the "
         "slide. If challenged on age: frameworks permit age when it does not disadvantage the "
         "elderly, and here the elderly are favoured. The young-applicant effect is genuinely "
         "risk-based, they default about 2.5 times more often, and it is documented and monitored. "
         "Gender had no such defence, so it is gone, and it cost nothing.")

# ============================================================== SLIDE 11 ===
s = new_slide(11)
kicker(s, "Live Demo")
headline(s, "The underwriting desk, running on this laptop", top=0.95, size=27)
card(s, 0.7, 1.85, 11.9, 2.35, fill=CARD_TINT)
accent_edge(s, 0.7, 1.85, 2.35, BLUE)
body_text(s, "I am about to open a small web app, type in an applicant, and let the real model "
             "decide in front of you.", 1.05, 2.1, 11.2, 0.6, size=17, color=NAVY, bold=True)
cols = [("You will see the decision", "the probability, and approve or decline at the 7% line"),
        ("You will see the money", "expected loss against expected margin, for that exact loan"),
        ("You will see the reasons", "the same four that would go on the customer's letter")]
cx = 1.05
for title, sub in cols:
    body_text(s, title, cx, 2.85, 3.5, 0.35, size=13.5, color=BLUE, bold=True)
    body_text(s, sub, cx, 3.2, 3.5, 0.75, size=12, color=MUTED)
    cx += 3.75
takeaway(s, 0.7, 4.5, 11.9, 1.05,
         "Nothing here is pre-recorded. It is the same model file I validated, scoring an applicant "
         "you choose, at the moment you choose it.", label="WHY I AM DOING THIS LIVE", size=13)
body_text(s, "If the demo misbehaves, the decline letter on slide 8 is the same output from the "
             "same model, and I will simply talk you through that instead.",
          0.7, 5.75, 11.9, 0.5, size=11.5, color=FAINT, italic=True)
notes(s, "Have the app already running at 127.0.0.1:5000 before you reach this slide. Click Thin "
         "file, talk while it scores, then read the money row out loud: expected loss against "
         "expected margin, and the net. Then point at the reasons and say these are the same four "
         "that go on the letter. Sixty seconds, then move on. Do not troubleshoot live.")

# ============================================================== SLIDE 12 ===
s = new_slide(12)
kicker(s, "How The Demo Works")
headline(s, "The model decides. The AI only writes it up.", top=0.95, size=27)
boxes = [
    ("1", "You type\nthe applicant", "in the browser", BLUE),
    ("2", "The trained\nmodel scores it", "probability of default", BLUE),
    ("3", "The maths\nexplains it", "which factors, how much", BLUE),
    ("4", "The AI\nnarrates it", "plain English paragraph", ORANGE),
]
bw, bgap = 2.62, 0.52
x, top, h = 0.75, 2.1, 1.95
for num, title, sub, col in boxes:
    card(s, x, top, bw, h, fill=ORANGE_TINT if col == ORANGE else CARD_TINT,
         line_color=ORANGE if col == ORANGE else BORDER)
    accent_edge(s, x, top, h, col)
    body_text(s, num, x + 0.22, top + 0.14, 1.0, 0.35, size=11, color=col, bold=True)
    body_text(s, title, x + 0.22, top + 0.48, bw - 0.45, 0.9, size=14.5, color=NAVY, bold=True)
    body_text(s, sub, x + 0.22, top + 1.35, bw - 0.45, 0.45, size=11.5, color=MUTED)
    if num != "4":
        chevron(s, x + bw + (bgap - 0.3) / 2, top + h / 2 - 0.14, 0.3, color=BLUE)
    x += bw + bgap

card(s, 0.7, 4.5, 11.9, 1.5, fill=NAVY, line_color=NAVY)
body_text(s, "The AI never makes the decision, and it cannot change one.",
          1.0, 4.68, 11.3, 0.45, size=18, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
body_text(s, "It receives a decision that has already been made, along with the numbers behind it, "
             "and turns them into a paragraph a customer or a committee can read. Unplug it and the "
             "decision, the money and the four reasons are identical.",
          1.0, 5.15, 11.3, 0.75, size=13, color=RGBColor(0xDC, 0xE7, 0xF2))
takeaway(s, 0.7, 6.2, 11.9, 0.85,
         "That separation is deliberate: everything auditable stays deterministic, and the only "
         "non-deterministic part is the prose.", label="WHY IT IS BUILT THIS WAY", size=12.5)
notes(s, "This is the slide a risk officer will care about most in the demo section. The governance "
         "point: the language model sits strictly downstream of the decision. It is given the "
         "probability, the decision and the SHAP factors as facts, and instructed to use only "
         "those numbers. If the API key is absent or the network is down, the app falls back to a "
         "deterministic template and the committee sees the same decision and the same reasons.")

# ============================================================== SLIDE 13 ===
s = new_slide(13)
tb = s.shapes.add_textbox(Inches(0), Inches(2.4), SW, Inches(0.4))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "THANK YOU"
r.font.name = BODY_FONT; r.font.bold = True; r.font.size = Pt(13); r.font.color.rgb = ORANGE
try:
    r._r.get_or_add_rPr().set('spc', '300')
except Exception:
    pass
tb2 = s.shapes.add_textbox(Inches(0), Inches(2.85), SW, Inches(1.3))
p2 = tb2.text_frame.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "Questions"
r2.font.name = HEAD_FONT; r2.font.bold = True; r2.font.size = Pt(54); r2.font.color.rgb = NAVY
body_text(s, "Kaustubh Patil  ·  Credit Scoring Capstone  ·  Model Review Committee",
          0, 4.25, 13.333, 0.4, size=13, color=MUTED, align=PP_ALIGN.CENTER)

summary = [("+$578.7M", "net benefit", GREEN), ("0.774", "test AUROC", BLUE),
           ("7%", "cut-off, where a loan breaks even", NAVY),
           ("Deploy", "with conditions", ORANGE)]
w, gap = 2.75, 0.25
x = (13.333 - (len(summary) * w + (len(summary) - 1) * gap)) / 2
for val, lab, col in summary:
    card(s, x, 5.3, w, 1.15, fill=CARD_TINT)
    accent_edge(s, x, 5.3, 1.15, col)
    big_num(s, val, x + 0.2, 5.42, w - 0.4, size=21, color=col)
    body_text(s, lab, x + 0.2, 5.95, w - 0.4, 0.4, size=11, color=MUTED)
    x += w + gap
notes(s, "Stop talking and open the floor. If nothing comes immediately, offer to score another "
         "applicant in the web app, or to walk back through the profit curve and what moving the "
         "threshold would cost.")

prs.save("reports/Credit_Scoring_Model_Review.pptx")
print(f"Saved reports/Credit_Scoring_Model_Review.pptx with {len(prs.slides._sldIdLst)} slides")
